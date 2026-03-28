"""
Генерация сравнительного Excel-отчёта по двум магазинам.
"""

import os
import re
import logging
from datetime import datetime

import pytz
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Side

from bot.config import THRESHOLD_A, THRESHOLD_B, REPORTS_DIR
from bot.calculations import aggregate_by_article
from bot.excel_styles import (
    HEADER_BG, HEADER_FG, GROUP_COLORS, WB_CABINET_URL,
    CMP_RAISE_BG, CMP_LOWER_BG, CMP_PAIR_ALT_BG,
    thin_border, fill,
    apply_header_style, write_legend_block, make_group_border,
)

logger = logging.getLogger(__name__)


def generate_comparison_report(
    store1_data: list[dict],
    store2_data: list[dict],
    store1_name: str,
    store2_name: str,
    days_threshold: int = 7,
    threshold_a: float = THRESHOLD_A,
    threshold_b: float = THRESHOLD_B,
) -> str | None:
    """
    Генерирует сравнительный Excel-отчёт по двум магазинам.

    Сопоставляет товары по supplier_article. Для каждой пары показывает
    данные обоих магазинов и рекомендацию по цене.

    Returns:
        Путь к файлу или None если нет совпадающих позиций.
    """
    logger.info(f"=== Сравнительный отчёт: {store1_name} vs {store2_name} ===")

    agg1 = aggregate_by_article(store1_data)
    agg2 = aggregate_by_article(store2_data)

    # Общие артикулы
    common_articles = sorted(set(agg1.keys()) & set(agg2.keys()))
    only_s1 = len(set(agg1.keys()) - set(agg2.keys()))
    only_s2 = len(set(agg2.keys()) - set(agg1.keys()))

    logger.info(f"Совпадений: {len(common_articles)}, только в маг.1: {only_s1}, только в маг.2: {only_s2}")

    if not common_articles:
        return None

    # Создаём Excel вручную (для merged cells)
    wb = Workbook()
    ws = wb.active
    ws.title = "Сравнение"

    # Шапка
    headers = ['Артикул', 'Магазин', 'ID (WB)', 'Остаток', 'Дней осталось',
               'Продаж/день', 'Группа', 'Цена ₽', 'Рекомендация']
    col_widths = {1: 22, 2: 24, 3: 15, 4: 12, 5: 16, 6: 15, 7: 11, 8: 13, 9: 18}

    apply_header_style(ws, col_widths, headers=headers)

    # Стили данных
    data_font = Font(name="Arial", size=10, color="1A1A2E")
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    hyperlink_font = Font(name="Arial", size=10, color="1155CC", underline="single")
    inner_thin = Side(style="thin", color="CCCCCC")
    pair_thick = Side(style="medium", color="888888")

    num_cols = len(headers)
    current_row = 2

    for pair_idx, article in enumerate(common_articles):
        s1 = agg1[article]
        s2 = agg2[article]

        # Определяем рекомендацию
        avg1 = s1.get('avg_per_day') or 0
        avg2 = s2.get('avg_per_day') or 0
        max_avg = max(avg1, avg2)

        if max_avg == 0:
            rec1, rec2 = "—", "—"
        elif abs(avg1 - avg2) / max_avg < 0.1:
            rec1, rec2 = "—", "—"
        elif avg1 > avg2:
            rec1, rec2 = "↑ Повысить", "↓ Понизить"
        else:
            rec1, rec2 = "↓ Понизить", "↑ Повысить"

        row1 = current_row
        row2 = current_row + 1

        # Чередование фона пар
        pair_bg = CMP_PAIR_ALT_BG if pair_idx % 2 == 1 else None

        for row_num, store_data, store_name_val, rec, is_top in [
            (row1, s1, store1_name, rec1, True),
            (row2, s2, store2_name, rec2, False),
        ]:
            is_bottom = not is_top
            ws.row_dimensions[row_num].height = 18

            def _border(col_idx):
                return make_group_border(inner_thin, pair_thick, is_top, is_bottom, col_idx, num_cols)

            # B: Магазин
            c = ws.cell(row=row_num, column=2, value=store_name_val)
            c.font = data_font
            c.alignment = left
            c.border = _border(2)

            # C: ID (WB) — гиперссылка
            nm_id = store_data.get('nm_id')
            c = ws.cell(row=row_num, column=3, value=nm_id)
            if nm_id:
                c.hyperlink = WB_CABINET_URL.format(nm_id)
                c.font = hyperlink_font
            else:
                c.font = data_font
            c.alignment = center
            c.border = _border(3)

            # D: Остаток
            c = ws.cell(row=row_num, column=4, value=store_data.get('stock_qty', 0))
            c.font = data_font
            c.alignment = center
            c.border = _border(4)
            c.number_format = "0"

            # E: Дней осталось
            dr = store_data.get('days_remaining')
            c = ws.cell(row=row_num, column=5, value=dr)
            c.font = data_font
            c.alignment = center
            c.border = _border(5)
            if dr is not None:
                c.number_format = "0.0"

            # F: Продаж/день
            c = ws.cell(row=row_num, column=6, value=store_data.get('avg_per_day', 0))
            c.font = data_font
            c.alignment = center
            c.border = _border(6)
            c.number_format = "0.00"

            # G: Группа
            group = store_data.get('product_group', '')
            c = ws.cell(row=row_num, column=7, value=group)
            c.font = Font(name="Arial", size=10, bold=True, color="1A1A2E")
            c.alignment = center
            c.border = _border(7)
            if group in GROUP_COLORS:
                c.fill = fill(GROUP_COLORS[group])

            # H: Цена
            price = store_data.get('price')
            c = ws.cell(row=row_num, column=8, value=price)
            c.font = data_font
            c.alignment = center
            c.border = _border(8)
            if price is not None:
                c.number_format = "0.00"

            # I: Рекомендация
            c = ws.cell(row=row_num, column=9, value=rec)
            c.font = Font(name="Arial", size=10, bold=True, color="1A1A2E")
            c.alignment = center
            c.border = _border(9)
            if "Повысить" in rec:
                c.fill = fill(CMP_RAISE_BG)
            elif "Понизить" in rec:
                c.fill = fill(CMP_LOWER_BG)

            # Фон чередования
            if pair_bg:
                for col in range(2, num_cols + 1):
                    existing = ws.cell(row=row_num, column=col)
                    if existing.fill == PatternFill():
                        existing.fill = fill(pair_bg)

        # A: Артикул (merged) с толстой рамкой
        ws.merge_cells(start_row=row1, start_column=1, end_row=row2, end_column=1)
        c = ws.cell(row=row1, column=1, value=article)
        c.font = Font(name="Arial", size=10, bold=True, color="1A1A2E")
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c.border = make_group_border(inner_thin, pair_thick, True, True, 1, num_cols)
        ws.cell(row=row2, column=1).border = make_group_border(
            inner_thin, pair_thick, False, True, 1, num_cols,
        )

        current_row += 2

    # Итоги внизу
    summary_row = current_row + 1
    summary_font = Font(name="Arial", size=9, bold=True, color="555555")

    ws.cell(row=summary_row, column=1, value=f"Совпадающих позиций: {len(common_articles)}").font = summary_font
    ws.cell(row=summary_row + 1, column=1, value=f"Только в {store1_name}: {only_s1}").font = summary_font
    ws.cell(row=summary_row + 2, column=1, value=f"Только в {store2_name}: {only_s2}").font = summary_font

    # Легенда
    legend_row = summary_row + 4
    write_legend_block(ws, legend_row, [
        ("— Группы товаров —", "title"),
        (f"A: ходовые (≥{threshold_a} шт/день) → среднее по 7 дням", "item"),
        (f"B: средние (≥{threshold_b} шт/день) → среднее по 14 дням", "item"),
        (f"C: редкие (<{threshold_b} шт/день) → среднее по 30 дням", "item"),
        ("", "item"),
        ("— Рекомендации —", "title"),
        ("↑ Повысить — товар продаётся лучше конкурента (можно повысить цену)", "item"),
        ("↓ Понизить — товар отстаёт (снизить цену для стимуляции спроса)", "item"),
        ("— паритет (разница продаж < 10%)", "item"),
    ])

    # Сохранение
    def _short_name(full_name):
        if not full_name:
            return "store"
        short = full_name.split('(')[0].strip()
        if not short:
            short = full_name.split()[0] if full_name.split() else "store"
        return re.sub(r'[^\w-]', '', short).strip()[:20]

    os.makedirs(REPORTS_DIR, exist_ok=True)
    timestamp = datetime.now(pytz.timezone('Europe/Moscow')).strftime('%Y%m%d_%H%M')
    output_path = os.path.join(
        REPORTS_DIR,
        f'comparison_{_short_name(store1_name)}_{_short_name(store2_name)}_{timestamp}.xlsx',
    )

    wb.save(output_path)
    logger.info(f"✓ Сравнительный отчёт: {output_path}")
    return output_path
