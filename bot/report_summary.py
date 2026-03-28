"""
Генерация сводного Excel-отчёта по всем магазинам.
"""

import os
import logging
from datetime import datetime

import pytz
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Side
from openpyxl.comments import Comment

from bot.config import THRESHOLD_A, THRESHOLD_B, REPORTS_DIR
from bot.calculations import aggregate_by_article, merge_wh_by_name
from bot.excel_styles import (
    GROUP_COLORS, WB_CABINET_URL, SUMMARY_PAIR_ALT_BG,
    fill,
    apply_header_style, write_legend_block, make_group_border,
)

logger = logging.getLogger(__name__)


def generate_summary_report(
    all_stores_data: dict,
    all_warehouse_data: dict,
    days_threshold: int = 7,
    threshold_a: float = THRESHOLD_A,
    threshold_b: float = THRESHOLD_B,
) -> str | None:
    """
    Генерирует сводный Excel-отчёт по всем магазинам.

    Показывает только товары (по supplier_article), присутствующие в 2+ магазинах.
    Вместо рекомендации — детализация остатков по складам.

    Args:
        all_stores_data: {store_name: [product_rows]}
        all_warehouse_data: {store_name: [{nm_id, warehouse_name, quantity}]}

    Returns:
        Путь к файлу или None если нет совпадающих позиций.
    """
    logger.info(f"=== Сводный отчёт: {len(all_stores_data)} магазинов ===")

    # Агрегация по артикулу для каждого магазина
    aggregated = {}
    for store_name, rows in all_stores_data.items():
        aggregated[store_name] = aggregate_by_article(rows)

    # Маппинг nm_id → supplier_article из product data
    nm_to_article = {}
    for store_name, rows in all_stores_data.items():
        for r in rows:
            nm_to_article.setdefault(store_name, {})[r['nm_id']] = r.get('supplier_article')

    # Дополняем маппинг из warehouse data (для nm_id без заказов)
    for store_name, wh_rows in all_warehouse_data.items():
        store_nm_map = nm_to_article.setdefault(store_name, {})
        for wr in wh_rows:
            if wr['nm_id'] not in store_nm_map:
                sa = wr.get('supplier_article')
                if sa:
                    store_nm_map[wr['nm_id']] = sa

    # Построение индекса складов по article (не nm_id)
    wh_index = {}  # {store_name: {article: [{warehouse_name, quantity, in_way_from_client}]}}
    for store_name, wh_rows in all_warehouse_data.items():
        store_wh = {}
        store_nm_map = nm_to_article.get(store_name, {})
        for wr in wh_rows:
            article = store_nm_map.get(wr['nm_id'])
            if not article:
                continue
            store_wh.setdefault(article, []).append({
                'warehouse_name': wr['warehouse_name'],
                'quantity': wr['quantity'],
                'in_way_from_client': wr.get('in_way_from_client', 0),
            })
        wh_index[store_name] = store_wh

    # Найти артикулы, присутствующие в 2+ магазинах
    article_stores = {}  # {article: [store_name, ...]}
    for store_name, agg in aggregated.items():
        for article in agg:
            article_stores.setdefault(article, []).append(store_name)

    common_articles = sorted([
        art for art, stores in article_stores.items() if len(stores) >= 2
    ])

    logger.info(f"Артикулов в 2+ магазинах: {len(common_articles)}")

    if not common_articles:
        return None

    # Создаём Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "Сводный отчёт"

    # Шапка
    headers = ['Артикул', 'Магазин', 'ID (WB)', 'Остаток', 'Дней осталось',
               'Продаж/день', 'Группа', 'Цена ₽', 'Остатки по складам']
    col_widths = {1: 22, 2: 24, 3: 15, 4: 12, 5: 16, 6: 15, 7: 11, 8: 13, 9: 36}

    apply_header_style(ws, col_widths, headers=headers)

    # Стили данных
    data_font = Font(name="Arial", size=10, color="1A1A2E")
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    wrap_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    hyperlink_font = Font(name="Arial", size=10, color="1155CC", underline="single")
    inner_thin = Side(style="thin", color="CCCCCC")
    group_thick = Side(style="medium", color="888888")

    num_cols = len(headers)
    current_row = 2

    for group_idx, article in enumerate(common_articles):
        stores_with_article = article_stores[article]
        group_size = len(stores_with_article)
        group_start = current_row

        # Чередование фона
        group_bg = SUMMARY_PAIR_ALT_BG if group_idx % 2 == 1 else None

        for store_idx, store_name in enumerate(stores_with_article):
            row_num = current_row
            is_top = (store_idx == 0)
            is_bottom = (store_idx == group_size - 1)
            store_data = aggregated[store_name][article]

            def _border(col_idx):
                return make_group_border(inner_thin, group_thick, is_top, is_bottom, col_idx, num_cols)

            # B: Магазин
            c = ws.cell(row=row_num, column=2, value=store_name)
            c.font = data_font
            c.alignment = left
            c.border = _border(2)

            # C: ID (WB) — гиперссылка
            nm_id = store_data.get('nm_id')
            c = ws.cell(row=row_num, column=3, value=nm_id)
            if nm_id:
                c.hyperlink = WB_CABINET_URL.format(nm_id)
                c.font = hyperlink_font
            else:
                c.font = data_font
            c.alignment = center
            c.border = _border(3)

            # Warehouse data (один раз для D, комментария и I)
            wh_list = wh_index.get(store_name, {}).get(article, [])
            wh_list = merge_wh_by_name(wh_list)
            total_iwfc = sum(w.get('in_way_from_client', 0) for w in wh_list)
            wh_list = [w for w in wh_list if w['quantity'] > 0]
            wh_list.sort(key=lambda w: w['quantity'], reverse=True)

            # D: Остаток (quantity — товар на складе, возвраты отдельно)
            if wh_list:
                wh_total_qty = sum(w['quantity'] for w in wh_list)
                display_stock = f"{wh_total_qty} (+{total_iwfc})" if total_iwfc > 0 else str(wh_total_qty)
            else:
                stock_val = store_data.get('stock_qty', 0)
                iwfc_total = store_data.get('in_way_from_client', 0)
                display_stock = f"{stock_val} (+{iwfc_total})" if iwfc_total > 0 else str(stock_val)
            c = ws.cell(row=row_num, column=4, value=display_stock)
            c.font = data_font
            c.alignment = center
            c.border = _border(4)

            # Комментарий к остатку: детализация по складам
            if wh_list or total_iwfc > 0:
                detailed_lines = [f"{w['warehouse_name']}: {w['quantity']} шт" for w in wh_list]
                detailed = "\n".join(detailed_lines) if detailed_lines else "нет"
                wh_total_qty = sum(w['quantity'] for w in wh_list)
                comment_text = (
                    f"Остатки по складам:\n{detailed}\n\n"
                    f"Итого на складах: {wh_total_qty} шт"
                )
                if total_iwfc > 0:
                    comment_text += f"\nВ возврате: {total_iwfc} шт"
                comment = Comment(comment_text, "WB Analiz")
                comment.width = 300
                total_lines = len(wh_list) + 8
                comment.height = total_lines * 14
                c.comment = comment

            # E: Дней осталось
            dr = store_data.get('days_remaining')
            c = ws.cell(row=row_num, column=5, value=dr)
            c.font = data_font
            c.alignment = center
            c.border = _border(5)
            if dr is not None:
                c.number_format = "0.0"

            # F: Продаж/день
            c = ws.cell(row=row_num, column=6, value=store_data.get('avg_per_day', 0))
            c.font = data_font
            c.alignment = center
            c.border = _border(6)
            c.number_format = "0.00"

            # G: Группа
            group = store_data.get('product_group', '')
            c = ws.cell(row=row_num, column=7, value=group)
            c.font = Font(name="Arial", size=10, bold=True, color="1A1A2E")
            c.alignment = center
            c.border = _border(7)
            if group in GROUP_COLORS:
                c.fill = fill(GROUP_COLORS[group])

            # H: Цена
            price = store_data.get('price')
            c = ws.cell(row=row_num, column=8, value=price)
            c.font = data_font
            c.alignment = center
            c.border = _border(8)
            if price is not None:
                c.number_format = "0.00"

            # I: Остатки по складам — компактная строка
            if wh_list:
                compact_parts = [f"{w['warehouse_name']}: {w['quantity']}" for w in wh_list]
                if total_iwfc > 0:
                    compact_parts.append(f"+{total_iwfc} возвр.")
                compact = " | ".join(compact_parts)
            else:
                stock_val = store_data.get('stock_qty', 0)
                iwfc_total = store_data.get('in_way_from_client', 0)
                parts = []
                if stock_val > 0:
                    parts.append(f"итого: {stock_val}")
                if iwfc_total > 0:
                    parts.append(f"+{iwfc_total} возвр.")
                compact = " | ".join(parts) if parts else "—"

            c = ws.cell(row=row_num, column=9, value=compact)
            c.font = Font(name="Arial", size=9, color="555555")
            c.alignment = wrap_left
            c.border = _border(9)

            # Фон чередования
            if group_bg:
                for col in range(2, num_cols + 1):
                    existing = ws.cell(row=row_num, column=col)
                    if existing.fill == PatternFill():
                        existing.fill = fill(group_bg)

            current_row += 1

        # A: Артикул (merged) с толстой рамкой
        if group_size > 1:
            ws.merge_cells(
                start_row=group_start, start_column=1,
                end_row=group_start + group_size - 1, end_column=1,
            )
        c = ws.cell(row=group_start, column=1, value=article)
        c.font = Font(name="Arial", size=10, bold=True, color="1A1A2E")
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c.border = make_group_border(inner_thin, group_thick, True, True, 1, num_cols)
        # Нижняя ячейка merged-области
        if group_size > 1:
            ws.cell(row=group_start + group_size - 1, column=1).border = make_group_border(
                inner_thin, group_thick, False, True, 1, num_cols,
            )

    # Итоги
    summary_row = current_row + 1
    summary_font = Font(name="Arial", size=9, bold=True, color="555555")
    total_articles = len(common_articles)
    total_stores = len(all_stores_data)

    ws.cell(row=summary_row, column=1,
            value=f"Общих позиций (в 2+ магазинах): {total_articles}").font = summary_font
    ws.cell(row=summary_row + 1, column=1,
            value=f"Магазинов в отчёте: {total_stores}").font = summary_font

    # Легенда
    legend_row = summary_row + 3
    write_legend_block(ws, legend_row, [
        ("— Группы товаров —", "title"),
        (f"A: ходовые (≥{threshold_a} шт/день) → среднее по 7 дням", "item"),
        (f"B: средние (≥{threshold_b} шт/день) → среднее по 14 дням", "item"),
        (f"C: редкие (<{threshold_b} шт/день) → среднее по 30 дням", "item"),
        ("", "item"),
        ("— Остатки по складам —", "title"),
        ("Наведите мышь на ячейку для детализации по складам", "item"),
    ])

    # Сохранение
    os.makedirs(REPORTS_DIR, exist_ok=True)
    timestamp = datetime.now(pytz.timezone('Europe/Moscow')).strftime('%Y%m%d_%H%M')
    output_path = os.path.join(REPORTS_DIR, f'summary_{timestamp}.xlsx')

    wb.save(output_path)
    logger.info(f"✓ Сводный отчёт: {output_path}")
    return output_path
